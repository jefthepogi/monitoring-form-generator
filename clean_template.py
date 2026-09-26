from docx import Document

def clean_word_xml_fragmentation(input_path="template.docx", output_path="template_clean.docx"):
    doc = Document(input_path)
    
    # 1. Clean standard paragraphs (fixes the ghost text in Title/Degree/Term)
    for p in doc.paragraphs:
        if "{{" in p.text or "{%" in p.text:
            # Reassigning the text destroys hidden XML fragments and merges it into one block
            p.text = p.text 

    # 2. Clean table cells (fixes the broken {% tr %} tags)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{" in p.text or "{%" in p.text:
                        p.text = p.text
                        
    doc.save(output_path)
    print(f"Success! Cleaned template saved as: {output_path}")

if __name__ == "__main__":
    clean_word_xml_fragmentation()